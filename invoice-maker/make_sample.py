"""動作確認用の売上明細（sample/売上明細.xlsx）を作成する。取引先・金額はすべて架空です。"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook

ROWS = [
    ("2026/09/02", "株式会社ミナト食品", "業務用だしパック", 20, 680, 10),
    ("2026/09/02", "株式会社ミナト食品", "焼き菓子詰め合わせ", 12, 1250, 8),
    ("2026/09/15", "株式会社ミナト食品", "配送料", 1, 1500, 10),
    ("2026/09/05", "カフェ ひだまり", "コーヒー豆 1kg", 3, 3980, 8),
    ("2026/09/05", "カフェ ひだまり", "紙カップ 100個", 5, 890, 10),
    ("2026/09/21", "カフェ ひだまり", "コーヒー豆 1kg", 2, 3980, 8),
    ("2026/09/10", "あおば設計事務所", "図面データ整理", 1, 45000, 10),
    ("2026/09/18", "あおば設計事務所", "出張費", 1, 3333, 10),
    ("2026/09/25", "カフェ ひだまり", "シロップ", "二本", 450, 8),  # 数量が数値でない行（エラー例）
    ("2026/10/01", "株式会社ミナト食品", "業務用だしパック", 10, 680, 10),
]


def main() -> None:
    out = Path(__file__).parent / "sample" / "売上明細.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "売上明細"
    ws.append(["日付", "取引先", "品目", "数量", "単価", "税率"])
    for date, *rest in ROWS:
        ws.append([dt.datetime.strptime(date, "%Y/%m/%d"), *rest[:-1], f"{rest[-1]}%"])
        ws.cell(ws.max_row, 1).number_format = "yyyy/mm/dd"
    for col, width in zip("ABCDEF", (12, 22, 22, 8, 10, 8)):
        ws.column_dimensions[col].width = width
    wb.save(out)
    print(f"作成しました: {out}")


if __name__ == "__main__":
    main()
