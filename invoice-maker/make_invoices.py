"""売上明細（Excel）から取引先ごとの請求書（Excel）を一括作成するツール。

- 取引先ごと・対象月ごとに明細をまとめて請求書を1ファイルずつ作成
- 標準税率10%／軽減税率8%の混在に対応（インボイス制度に合わせ、税率ごとに1回だけ端数処理）
- 作成結果を「請求一覧.xlsx」にまとめて出力
"""
from __future__ import annotations

import argparse
import datetime as dt
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import ROUND_DOWN, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REQUIRED_COLUMNS = ["日付", "取引先", "品目", "数量", "単価", "税率"]


@dataclass
class Line:
    date: dt.date
    client: str
    item: str
    qty: Decimal
    price: Decimal
    rate: int  # 10 or 8

    @property
    def amount(self) -> Decimal:
        return self.qty * self.price


@dataclass
class Invoice:
    client: str
    month: str  # YYYY-MM
    lines: list[Line] = field(default_factory=list)

    def subtotal_by_rate(self) -> dict[int, Decimal]:
        totals: dict[int, Decimal] = defaultdict(Decimal)
        for line in self.lines:
            totals[line.rate] += line.amount
        return dict(sorted(totals.items(), reverse=True))

    def tax_by_rate(self) -> dict[int, int]:
        """税率ごとに合計してから1回だけ切り捨てる（インボイス制度の端数処理ルール）。"""
        return {
            rate: int((subtotal * rate / 100).to_integral_value(rounding=ROUND_DOWN))
            for rate, subtotal in self.subtotal_by_rate().items()
        }

    @property
    def subtotal(self) -> int:
        return int(sum(self.subtotal_by_rate().values()))

    @property
    def tax(self) -> int:
        return sum(self.tax_by_rate().values())

    @property
    def total(self) -> int:
        return self.subtotal + self.tax


class InputError(Exception):
    pass


def _to_date(value, row_no: int) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise InputError(f"{row_no}行目: 日付が読めません（{value!r}）")


def _to_decimal(value, name: str, row_no: int) -> Decimal:
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except Exception:
        raise InputError(f"{row_no}行目: {name}が数値ではありません（{value!r}）") from None
    if not number.is_finite():
        raise InputError(f"{row_no}行目: {name}が数値ではありません（{value!r}）")
    return number


def _to_rate(value, row_no: int) -> int:
    text = str(value).replace("%", "").strip()
    try:
        number = float(text)
    except ValueError:
        raise InputError(f"{row_no}行目: 税率が読めません（{value!r}）") from None
    if number < 1:  # 0.1 / 0.08 の形式
        number *= 100
    rate = int(round(number))
    if rate not in (8, 10):
        raise InputError(f"{row_no}行目: 税率は10%か8%にしてください（{value!r}）")
    return rate


def read_lines(path: Path) -> tuple[list[Line], list[str]]:
    """売上明細を読み込む。読めない行はエラーとして返し、処理は続ける。"""
    sheet = load_workbook(path, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise InputError(f"列が見つかりません: {', '.join(missing)}")
    index = {name: header.index(name) for name in REQUIRED_COLUMNS}

    lines: list[Line] = []
    errors: list[str] = []
    for row_no, row in enumerate(rows, start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        try:
            get = lambda name: row[index[name]]  # noqa: E731
            client = str(get("取引先") or "").strip()
            item = str(get("品目") or "").strip()
            if not client:
                raise InputError(f"{row_no}行目: 取引先が空です")
            if not item:
                raise InputError(f"{row_no}行目: 品目が空です")
            lines.append(Line(
                date=_to_date(get("日付"), row_no),
                client=client,
                item=item,
                qty=_to_decimal(get("数量"), "数量", row_no),
                price=_to_decimal(get("単価"), "単価", row_no),
                rate=_to_rate(get("税率"), row_no),
            ))
        except InputError as e:
            errors.append(str(e))
    return lines, errors


def group_invoices(lines: list[Line], month: str | None = None) -> list[Invoice]:
    invoices: dict[tuple[str, str], Invoice] = {}
    for line in sorted(lines, key=lambda l: (l.client, l.date)):
        key_month = line.date.strftime("%Y-%m")
        if month and key_month != month:
            continue
        key = (line.client, key_month)
        invoices.setdefault(key, Invoice(client=line.client, month=key_month)).lines.append(line)
    return list(invoices.values())


def due_date(month: str) -> dt.date:
    """翌月末を支払期限にする。"""
    year, mon = map(int, month.split("-"))
    mon += 2  # 翌々月の1日から1日戻すと翌月末になる
    year += (mon - 1) // 12
    mon = (mon - 1) % 12 + 1
    return dt.date(year, mon, 1) - dt.timedelta(days=1)


THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E8EEF8")
YEN = '"¥"#,##0'


def write_invoice(inv: Invoice, issuer: dict, number: str, issue_date: dt.date, out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "請求書"
    for col, width in zip("ABCDEF", (12, 34, 8, 12, 8, 14)):
        ws.column_dimensions[col].width = width

    ws["A1"] = "請求書"
    ws["A1"].font = Font(size=18, bold=True)
    ws["E1"], ws["F1"] = "請求番号", number
    ws["E2"], ws["F2"] = "発行日", issue_date
    ws["F2"].number_format = "yyyy/mm/dd"

    ws["A3"] = f"{inv.client} 御中"
    ws["A3"].font = Font(size=13, bold=True)
    ws["A5"] = f"{inv.month.replace('-', '年')}月分のご請求です。"

    ws["D4"] = issuer["name"]
    ws["D5"] = f"登録番号 {issuer['registration']}"
    ws["D6"] = issuer["address"]

    ws["A7"], ws["B7"] = "ご請求金額", inv.total
    ws["A7"].font = ws["B7"].font = Font(size=14, bold=True)
    ws["B7"].number_format = YEN
    ws["A8"], ws["B8"] = "お支払期限", due_date(inv.month)
    ws["B8"].number_format = "yyyy/mm/dd"
    ws["B8"].alignment = Alignment(horizontal="left")

    row = 10
    for col, title in zip("ABCDEF", ("日付", "品目", "数量", "単価", "税率", "金額")):
        cell = ws[f"{col}{row}"]
        cell.value, cell.fill, cell.border, cell.font = title, HEAD_FILL, BOX, Font(bold=True)
    for line in inv.lines:
        row += 1
        mark = "※" if line.rate == 8 else ""
        values = (line.date, f"{line.item}{mark}", float(line.qty), float(line.price), f"{line.rate}%", float(line.amount))
        for col, value in zip("ABCDEF", values):
            ws[f"{col}{row}"] = value
            ws[f"{col}{row}"].border = BOX
        ws[f"A{row}"].number_format = "mm/dd"
        ws[f"D{row}"].number_format = ws[f"F{row}"].number_format = YEN

    row += 2
    for rate, subtotal in inv.subtotal_by_rate().items():
        ws[f"D{row}"], ws[f"F{row}"] = f"{rate}%対象", float(subtotal)
        ws[f"D{row + 1}"], ws[f"F{row + 1}"] = f"消費税({rate}%)", inv.tax_by_rate()[rate]
        ws[f"F{row}"].number_format = ws[f"F{row + 1}"].number_format = YEN
        row += 2
    ws[f"D{row}"], ws[f"F{row}"] = "合計", inv.total
    ws[f"D{row}"].font = ws[f"F{row}"].font = Font(bold=True)
    ws[f"F{row}"].number_format = YEN
    if any(l.rate == 8 for l in inv.lines):
        ws[f"A{row + 2}"] = "※は軽減税率（8%）対象です。"

    ws.print_area = f"A1:F{row + 2}"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)


def safe_filename(name: str) -> str:
    return "".join("_" if c in '\\/:*?"<>|' else c for c in name)


def run(src: Path, out_dir: Path, issuer: dict, month: str | None, issue_date: dt.date) -> dict:
    lines, errors = read_lines(src)
    invoices = group_invoices(lines, month)
    out_dir.mkdir(parents=True, exist_ok=True)

    summary = Workbook()
    ws = summary.active
    ws.title = "請求一覧"
    ws.append(["請求番号", "取引先", "対象月", "明細数", "税抜", "消費税", "合計", "ファイル"])
    for no, inv in enumerate(invoices, start=1):
        number = f"{inv.month.replace('-', '')}-{no:03d}"
        filename = f"請求書_{inv.month}_{safe_filename(inv.client)}.xlsx"
        write_invoice(inv, issuer, number, issue_date, out_dir / filename)
        ws.append([number, inv.client, inv.month, len(inv.lines), inv.subtotal, inv.tax, inv.total, filename])
    if errors:
        err = summary.create_sheet("読み込みエラー")
        err.append(["内容"])
        for message in errors:
            err.append([message])
    summary.save(out_dir / "請求一覧.xlsx")
    return {"lines": len(lines), "invoices": len(invoices), "errors": errors,
            "total": sum(inv.total for inv in invoices)}


def main() -> None:
    parser = argparse.ArgumentParser(description="売上明細から取引先ごとの請求書を作成します")
    parser.add_argument("source", type=Path, help="売上明細のExcelファイル")
    parser.add_argument("-o", "--out", type=Path, default=Path("output"), help="出力フォルダ")
    parser.add_argument("--month", help="対象月（例: 2026-09）。省略時はすべての月")
    parser.add_argument("--issuer", default="サンプル商事", help="請求元の名前")
    parser.add_argument("--registration", default="T0000000000000", help="インボイス登録番号")
    parser.add_argument("--address", default="東京都架空区1-2-3", help="請求元の住所")
    args = parser.parse_args()

    issuer = {"name": args.issuer, "registration": args.registration, "address": args.address}
    result = run(args.source, args.out, issuer, args.month, dt.date.today())
    print(f"明細 {result['lines']}件 → 請求書 {result['invoices']}件（合計 {result['total']:,}円）")
    for message in result["errors"]:
        print(f"  読み込みエラー: {message}")
    print(f"出力先: {args.out}")


if __name__ == "__main__":
    main()
