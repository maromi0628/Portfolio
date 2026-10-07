import datetime as dt
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from make_invoices import InputError, Invoice, Line, due_date, group_invoices, read_lines, run


def line(client="A社", date=dt.date(2026, 9, 1), qty="1", price="100", rate=10, item="品目"):
    return Line(date, client, item, Decimal(qty), Decimal(price), rate)


def write_book(path, rows, header=("日付", "取引先", "品目", "数量", "単価", "税率")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


def test_tax_is_rounded_once_per_rate():
    # 明細ごとに切り捨てると 9+9+9=27 円だが、税率ごとにまとめると 99*3*10% = 29 円
    inv = Invoice("A社", "2026-09", [line(price="99") for _ in range(3)])
    assert inv.tax_by_rate() == {10: 29}
    assert inv.total == 297 + 29


def test_mixed_rates():
    inv = Invoice("A社", "2026-09", [line(price="1000", rate=10), line(price="1000", rate=8)])
    assert inv.subtotal_by_rate() == {10: Decimal("1000"), 8: Decimal("1000")}
    assert inv.tax_by_rate() == {10: 100, 8: 80}
    assert inv.total == 2180


def test_group_by_client_and_month():
    lines = [line("B社"), line("A社"), line("A社", dt.date(2026, 10, 1)), line("A社", dt.date(2026, 9, 30))]
    invoices = group_invoices(lines)
    assert [(i.client, i.month, len(i.lines)) for i in invoices] == [
        ("A社", "2026-09", 2), ("A社", "2026-10", 1), ("B社", "2026-09", 1)]
    assert [i.client for i in group_invoices(lines, "2026-10")] == ["A社"]


@pytest.mark.parametrize("month, expected", [
    ("2026-09", dt.date(2026, 10, 31)),
    ("2026-11", dt.date(2026, 12, 31)),
    ("2026-12", dt.date(2027, 1, 31)),
    ("2027-01", dt.date(2027, 2, 28)),
])
def test_due_date_is_end_of_next_month(month, expected):
    assert due_date(month) == expected


def test_bad_rows_are_reported_and_skipped(tmp_path):
    src = write_book(tmp_path / "in.xlsx", [
        ("2026/09/01", "A社", "作業", 2, "1,500", "10%"),
        ("2026/09/02", "A社", "食品", 1, 500, 0.08),
        ("2026/09/03", "", "作業", 1, 100, 10),
        ("2026/09/04", "A社", "作業", "二", 100, 10),
        ("2026/09/05", "A社", "作業", 1, 100, 5),
        ("9月6日", "A社", "作業", 1, 100, 10),
        (None, None, None, None, None, None),
    ])
    lines, errors = read_lines(src)
    assert [(l.price, l.rate) for l in lines] == [(Decimal("1500"), 10), (Decimal("500"), 8)]
    assert len(errors) == 4
    assert errors[0].startswith("4行目: 取引先が空")


def test_missing_column(tmp_path):
    src = write_book(tmp_path / "in.xlsx", [], header=("日付", "取引先", "品目", "数量", "単価"))
    with pytest.raises(InputError, match="税率"):
        read_lines(src)


def test_run_writes_invoices_and_summary(tmp_path):
    src = write_book(tmp_path / "in.xlsx", [
        ("2026/09/01", "A社", "作業", 1, 10000, 10),
        ("2026/09/02", "B/社", "食品", 2, 1000, 8),
        ("2026/09/03", "B/社", "作業", 1, "x", 10),
    ])
    issuer = {"name": "テスト", "registration": "T0000000000000", "address": "架空"}
    result = run(src, tmp_path / "out", issuer, None, dt.date(2026, 10, 1))
    assert result["invoices"] == 2 and result["total"] == 11000 + 2160
    assert (tmp_path / "out" / "請求書_2026-09_B_社.xlsx").exists()
    summary = load_workbook(tmp_path / "out" / "請求一覧.xlsx")
    assert summary.sheetnames == ["請求一覧", "読み込みエラー"]
    rows = list(summary["請求一覧"].iter_rows(min_row=2, values_only=True))
    assert rows[0][:7] == ("202609-001", "A社", "2026-09", 1, 10000, 1000, 11000)
