"""企業リスト抽出ツール

Web 上の企業一覧ページから会社情報を取得し、
表記ゆれの正規化 → 重複削除 → 除外リスト照合 → Excel 出力 までを一括で行います。

使い方:
    # 動作確認（同梱の架空サンプルを使用）
    python extractor.py --from-dir sample --exclude sample/exclude.txt -o result.xlsx

    # 実サイト（利用規約・robots.txt で許可されている場合のみ）
    python extractor.py --start-url https://example.com/list?page=1 --delay 3 -o result.xlsx

サイトごとの HTML 構造は SELECTORS を書き換えて対応します。
"""
from __future__ import annotations

import argparse
import re
import sys
import time
import unicodedata
import urllib.robotparser
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# サイトの HTML 構造に合わせて変更する
SELECTORS = {
    "item": "li.company",
    "name": ".name",
    "representative": ".rep",
    "address": ".address",
    "tel": ".tel",
    "url": ".url a",
    "next": "a.next",
}

USER_AGENT = "list-extractor/1.0 (+contact: see README)"
CORP_TYPES = ["株式会社", "有限会社", "合同会社", "合資会社", "合名会社", "一般社団法人", "(株)", "(有)", "(同)"]


@dataclass
class Company:
    name: str
    representative: str = ""
    address: str = ""
    tel: str = ""
    url: str = ""
    source: str = ""
    notes: list[str] = field(default_factory=list)


# ---------- 正規化 ----------

def nfkc(text: str) -> str:
    """全角英数・記号を半角に、空白を1つにそろえる"""
    text = unicodedata.normalize("NFKC", text or "")
    return re.sub(r"\s+", " ", text).strip()


def normalize_name(name: str) -> str:
    """表示用の会社名: (株)→株式会社 などをそろえる"""
    s = nfkc(name)
    s = s.replace("(株)", "株式会社").replace("(有)", "有限会社").replace("(同)", "合同会社")
    s = re.sub(r"(株式会社|有限会社|合同会社)\s+", r"\1", s)
    s = re.sub(r"\s+(株式会社|有限会社|合同会社)", r"\1", s)
    return s


def name_key(name: str) -> str:
    """照合用のキー: 法人格・空白・記号を除いた会社名"""
    s = normalize_name(name)
    for t in CORP_TYPES:
        s = s.replace(nfkc(t), "")
    return re.sub(r"[\s・.,、。]", "", s).lower()


def normalize_tel(tel: str) -> str:
    s = nfkc(tel).replace("ー", "-").replace("−", "-").replace("‐", "-")
    digits = re.sub(r"\D", "", s)
    return s if 10 <= len(digits) <= 11 else ""


def tel_key(tel: str) -> str:
    return re.sub(r"\D", "", tel)


# ---------- 取得 ----------

def parse_page(html: str, base_url: str) -> tuple[list[Company], str | None]:
    soup = BeautifulSoup(html, "html.parser")
    items = []
    for el in soup.select(SELECTORS["item"]):
        def text(key: str) -> str:
            node = el.select_one(SELECTORS[key])
            return node.get_text(strip=True) if node else ""
        link = el.select_one(SELECTORS["url"])
        items.append(Company(
            name=text("name"), representative=text("representative"), address=text("address"),
            tel=text("tel"), url=link.get("href", "") if link else "", source=base_url,
        ))
    nxt = soup.select_one(SELECTORS["next"])
    return items, (urljoin(base_url, nxt["href"]) if nxt and nxt.get("href") else None)


def crawl_web(start_url: str, delay: float, max_pages: int) -> list[Company]:
    rp = urllib.robotparser.RobotFileParser()
    parts = urlparse(start_url)
    rp.set_url(f"{parts.scheme}://{parts.netloc}/robots.txt")
    try:
        rp.read()
    except Exception:
        print("robots.txt を取得できませんでした。規約を手動で確認してください。", file=sys.stderr)

    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    url, results, pages = start_url, [], 0
    while url and pages < max_pages:
        if not rp.can_fetch(USER_AGENT, url):
            print(f"robots.txt で禁止されているため停止: {url}", file=sys.stderr)
            break
        res = session.get(url, timeout=20)
        res.raise_for_status()
        res.encoding = res.apparent_encoding
        items, url = parse_page(res.text, res.url)
        results += items
        pages += 1
        print(f"  {pages}ページ目: {len(items)}件")
        if url:
            time.sleep(delay)  # サイトに負荷をかけない
    return results


def crawl_dir(directory: Path) -> list[Company]:
    results = []
    for path in sorted(directory.glob("*.html")):
        items, _ = parse_page(path.read_text(encoding="utf-8"), path.name)
        results += items
        print(f"  {path.name}: {len(items)}件")
    return results


# ---------- クレンジング ----------

def cleanse(raw: list[Company], exclude_names: list[str]) -> tuple[list[Company], dict]:
    stats = {"取得件数": len(raw), "重複削除": 0, "除外リスト該当": 0, "会社名なし": 0}
    exclude_keys = {name_key(n) for n in exclude_names}
    seen_names: dict[str, Company] = {}
    seen_tels: dict[str, Company] = {}
    out: list[Company] = []

    for c in raw:
        c.name = normalize_name(c.name)
        c.representative = nfkc(c.representative)
        c.address = nfkc(c.address)
        c.tel = normalize_tel(c.tel)
        c.url = nfkc(c.url)
        if not c.name:
            stats["会社名なし"] += 1
            continue
        if name_key(c.name) in exclude_keys:
            stats["除外リスト該当"] += 1
            continue
        nk, tk = name_key(c.name), tel_key(c.tel)
        dup = seen_names.get(nk) or (seen_tels.get(tk) if tk else None)
        if dup:
            stats["重複削除"] += 1
            # 片方にしかない情報は残す
            for f in ("representative", "address", "tel", "url"):
                if not getattr(dup, f) and getattr(c, f):
                    setattr(dup, f, getattr(c, f))
            continue
        if not c.tel:
            c.notes.append("電話番号なし/形式不正")
        if not c.url:
            c.notes.append("HPなし")
        seen_names[nk] = c
        if tk:
            seen_tels[tk] = c
        out.append(c)

    stats["納品件数"] = len(out)
    return out, stats


# ---------- 出力 ----------

def write_xlsx(rows: list[Company], stats: dict, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "企業リスト"
    header = ["No", "会社名", "代表者名", "住所", "電話番号", "ホームページURL", "備考"]
    ws.append(header)
    for i, c in enumerate(rows, 1):
        ws.append([i, c.name, c.representative, c.address, c.tel, c.url, " / ".join(c.notes)])
    fill = PatternFill("solid", fgColor="DDE7F7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for col, width in zip(range(1, len(header) + 1), [6, 30, 14, 34, 16, 34, 22]):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("処理結果")
    for k, v in stats.items():
        summary.append([k, v])
    summary.column_dimensions["A"].width = 18
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="企業リスト抽出ツール")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--start-url", help="一覧ページの URL（次ページは自動でたどる）")
    src.add_argument("--from-dir", type=Path, help="保存済み HTML のフォルダ")
    ap.add_argument("--exclude", type=Path, help="除外する企業名リスト（1行1社）")
    ap.add_argument("--delay", type=float, default=3.0, help="ページ取得の間隔（秒）")
    ap.add_argument("--max-pages", type=int, default=1000)
    ap.add_argument("-o", "--output", type=Path, default=Path("result.xlsx"))
    args = ap.parse_args(argv)

    print("取得中...")
    raw = crawl_dir(args.from_dir) if args.from_dir else crawl_web(args.start_url, args.delay, args.max_pages)

    exclude = []
    if args.exclude:
        exclude = [l.strip() for l in args.exclude.read_text(encoding="utf-8").splitlines()
                   if l.strip() and not l.startswith("#")]
    rows, stats = cleanse(raw, exclude)
    write_xlsx(rows, stats, args.output)

    print("\n処理結果")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    print(f"\n→ {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
